"""Context & Objectives slide — W11/D2 (mode-agnostic).

Two-column layout pulling from ``payload.metadata.brief`` /
``payload._engagement_title`` (left) and the writer's executive
insights or summary (right). Falls back gracefully when the
session metadata bag is empty.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from .._layout import (
    DEFAULT_MUTED,
    DEFAULT_PRIMARY,
    DEFAULT_SECONDARY,
    SLIDE_WIDTH_IN,
    add_blank_slide,
    add_horizontal_band,
    add_paragraph,
    add_textbox,
    parse_hex,
)
from ..._base import payload_get
from ...one_pager_renderer import _coerce_to_list, _stringify_item
from ._base import SlideBuilderBase, SlideResult
from ._registry import register_slide


def _brief_text(payload: Any) -> str:
    """The engagement brief — usually carried on ``session.metadata.brief``
    or ``session.query`` and injected via the W10 service layer onto the
    payload's ``_engagement_brief`` underscore key. We tolerate three
    spellings + the engagement title as the last-resort fallback."""
    for key in ("_engagement_brief", "engagement_brief", "brief"):
        v = payload_get(payload, key, default=None)
        if isinstance(v, str) and v.strip():
            return v.strip()
    md = payload_get(payload, "metadata", default={}) or {}
    if isinstance(md, dict):
        b = md.get("brief") or md.get("query")
        if isinstance(b, str) and b.strip():
            return b.strip()
    return str(payload_get(payload, "_engagement_title", default="Engagement context not specified."))


def _objectives(payload: Any) -> list[str]:
    """Top-3 objectives. First tries executive_summary.objectives,
    then executive_insights[].text, then a single-bullet summary
    fallback."""
    es = payload_get(payload, "executive_summary", default=None)
    if isinstance(es, dict):
        objs = _coerce_to_list(es.get("objectives") or es.get("top_3_objectives") or [])
        out = [_stringify_item(x) for x in objs if _stringify_item(x)]
        if out:
            return out[:3]
    ei = _coerce_to_list(payload_get(payload, "executive_insights", default=[]))
    insights = [_stringify_item(x) for x in ei if _stringify_item(x)]
    if insights:
        return insights[:3]
    summ = payload_get(payload, "summary", default="")
    if isinstance(summ, str) and summ.strip():
        return [summ.strip()]
    return []


@register_slide("context")
class ContextSlide(SlideBuilderBase):
    def build(
        self,
        presentation: Any,
        payload: Any,
        firm_branding: dict[str, Any],
        citations: list[Any],
        deck_context: Any = None,
    ) -> SlideResult:
        primary = (firm_branding or {}).get("primary_color") or DEFAULT_PRIMARY
        secondary_hex = (firm_branding or {}).get("secondary_color") or DEFAULT_SECONDARY

        slide = add_blank_slide(presentation)
        objectives = _objectives(payload)
        brief = _brief_text(payload)

        # Two-column layout (or single-column fallback).
        if objectives:
            col_width = (SLIDE_WIDTH_IN - 1.4) / 2 - 0.2
            col_top = 1.3
            col_height = SLIDE_WIDTH_IN  # height bound by textbox height arg below

            add_textbox(
                slide,
                left=0.5, top=col_top,
                width=col_width, height=0.4,
                text="Engagement context",
                font_size=14, bold=True,
                color=parse_hex(DEFAULT_MUTED),
                align=PP_ALIGN.LEFT,
            )
            brief_shape = add_textbox(
                slide,
                left=0.5, top=col_top + 0.5,
                width=col_width, height=5.5,
                text=brief[:1400],
                font_size=11,
                color=parse_hex(secondary_hex),
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            )

            obj_left = 0.5 + col_width + 0.4
            add_textbox(
                slide,
                left=obj_left, top=col_top,
                width=col_width, height=0.4,
                text="Objectives",
                font_size=14, bold=True,
                color=parse_hex(DEFAULT_MUTED),
                align=PP_ALIGN.LEFT,
            )
            obj_shape = add_textbox(
                slide,
                left=obj_left, top=col_top + 0.5,
                width=col_width, height=5.5,
                text="",
                font_size=12,
                color=parse_hex(secondary_hex),
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            )
            for obj in objectives:
                add_paragraph(
                    obj_shape.text_frame, str(obj)[:240],
                    font_size=12, bullet=True,
                    color=parse_hex(secondary_hex),
                )
        else:
            # Brief-only single column.
            add_textbox(
                slide,
                left=0.5, top=1.3,
                width=SLIDE_WIDTH_IN - 1.0, height=0.4,
                text="Engagement context",
                font_size=14, bold=True,
                color=parse_hex(DEFAULT_MUTED),
            )
            add_textbox(
                slide,
                left=0.5, top=1.8,
                width=SLIDE_WIDTH_IN - 1.0, height=5.0,
                text=brief[:2000],
                font_size=12,
                color=parse_hex(secondary_hex),
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            )

        return SlideResult(
            slide_index=len(presentation.slides) - 1,
            citation_ids=[],
        )
